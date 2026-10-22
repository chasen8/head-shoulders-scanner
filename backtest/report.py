"""Aggregate backtest trades into statistics + an Excel workbook (D40). Values only, no formulas.

Input:  results/<tag>/<market>/trades/*.csv, skips/*.csv, done/*.json
Output: results/<tag>/trades_all.csv (every trade), results/<tag>/summary.csv, results/<tag>/summary.md
        C:\\trading-data\\ihs-scanner-v2\\output\\回測結果_<date>.xlsx
"""
from __future__ import annotations

import glob
import json
import math
import time
from pathlib import Path

import numpy as np
import pandas as pd

from backtest import data as D
from backtest import engine as E

OUTPUT_DIR = Path(r"C:\trading-data\ihs-scanner-v2\output")
MIN_SAMPLE = 30
MIN_RISK_FOR_R = 0.005  # R statistics skip trades whose stop is < 0.5% away (a gap fill makes R meaningless there)


def _r(g: pd.DataFrame) -> pd.Series:
    return g.loc[g["risk_pct"] >= MIN_RISK_FOR_R, "R"]
MARKET_LABEL = {"crypto": "加密", "tw": "台股"}
FAMILY_LABEL = {"ihs": "頭肩底", "tri": "收斂三角"}
MODE_LABEL = {"A": "A 突破確認", "B": "B 右肩確認"}
STOP_LABEL = {"rs": "右肩低點", "inval": "偵測器失效價", "atr": "ATR 1.5×"}
EXIT_LABEL = {"TP1": "TP1 出場", "TP2": "TP2 出場", "half": "半半保本"}
DIR_LABEL = {1: "做多", -1: "做空（附錄）"}
TF_ORDER = ["4h", "1d", "1w", "全部"]


def _read_all(pattern: str) -> pd.DataFrame:
    frames = []
    for p in glob.glob(pattern):
        try:
            df = pd.read_csv(p, encoding="utf-8", low_memory=False)
        except pd.errors.EmptyDataError:
            continue
        if len(df):
            frames.append(df)
    return pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()


def load(tag: str) -> tuple[pd.DataFrame, pd.DataFrame, list[dict]]:
    base = D.RESULTS_DIR / tag
    trades = _read_all(str(base / "*" / "trades" / "*.csv"))
    skips = _read_all(str(base / "*" / "skips" / "*.csv"))
    done = []
    for p in glob.glob(str(base / "*" / "done" / "*.json")):
        with open(p, encoding="utf-8") as fh:
            info = json.load(fh)
        info["market"] = Path(p).parent.parent.name
        done.append(info)
    if len(trades):
        for c in ("entry_time", "exit_time"):
            trades[c] = pd.to_datetime(trades[c], utc=True, format="ISO8601")
        trades["year"] = trades["entry_time"].dt.year
        trades["variant"] = np.where(trades["family"] == "ihs", trades["mode"].fillna(""), trades["stop_kind"])
    return trades, skips, done


def max_losing_streak(rets: pd.Series) -> int:
    best = cur = 0
    for r in rets:
        cur = cur + 1 if r < 0 else 0
        best = max(best, cur)
    return best


def _p_two_prop(w1, n1, w2, n2) -> float | None:
    if min(n1, n2) == 0:
        return None
    p = (w1 + w2) / (n1 + n2)
    se = math.sqrt(p * (1 - p) * (1 / n1 + 1 / n2))
    if se == 0:
        return None
    z = abs(w1 / n1 - w2 / n2) / se
    return math.erfc(z / math.sqrt(2))


def _p_welch(a: pd.Series, b: pd.Series) -> float | None:
    if len(a) < 2 or len(b) < 2:
        return None
    se = math.sqrt(a.var(ddof=1) / len(a) + b.var(ddof=1) / len(b))
    if se == 0:
        return None
    z = abs(a.mean() - b.mean()) / se  # normal approximation (large samples)
    return math.erfc(z / math.sqrt(2))


def stats(g: pd.DataFrame) -> dict:
    g = g.sort_values(["entry_time", "symbol"])
    n = len(g)
    r = g["net_ret"]
    wins, losses = r[r > 0].sum(), -r[r < 0].sum()
    out = {
        "筆數": n,
        "標的數": g["symbol"].nunique(),
        "勝率%（停利出場比例）": 100 * g["tp_hit"].astype(bool).mean() if n else None,
        "正報酬比例%": 100 * (r > 0).mean() if n else None,
        "平均報酬%（每筆期望值）": 100 * r.mean() if n else None,
        "中位數報酬%": 100 * r.median() if n else None,
        "平均R（期望值R）": _r(g).mean() if n else None,
        "R中位數": _r(g).median() if n else None,
        "獲利因子": (wins / losses) if losses > 0 else None,
        "平均持有根數": g["bars_held"].mean() if n else None,
        "平均持有天數": g["days_held"].mean() if n else None,
        "最大連續虧損（筆）": max_losing_streak(r) if n else 0,
        "停損比例%": 100 * (g["exit_reason"] == "stop").mean() if n else None,
        "到期出場比例%": 100 * (g["exit_reason"] == "timeout").mean() if n else None,
        "樣本不足（<30）": "是" if n < MIN_SAMPLE else "",
    }
    if (g["exit"] == "half").all() and n:
        out["TP2達成%（半半）"] = 100 * g["tp2_hit"].astype(str).str.lower().eq("true").mean()
        out["保本出場%（半半）"] = 100 * (g["leg2_reason"] == "breakeven").mean()
    return out


GROUP_KEYS = ["market", "family", "direction", "variant", "timeframe", "exit"]


def summarize(tr: pd.DataFrame) -> pd.DataFrame:
    rows = []
    closed = tr[~tr["is_open"].astype(bool)]
    for keys, g in closed.groupby(["market", "family", "direction", "variant", "exit"], sort=True):
        for tf in TF_ORDER:
            sub = g if tf == "全部" else g[g["timeframe"] == tf]
            if not len(sub):
                continue
            m, fam, d, var, ex = keys
            rows.append({"市場": MARKET_LABEL[m], "型態": FAMILY_LABEL[fam], "方向": DIR_LABEL[int(d)],
                         "進場／停損": MODE_LABEL.get(var) if fam == "ihs" else STOP_LABEL[var],
                         "週期": tf, "出場策略": EXIT_LABEL[ex], **stats(sub),
                         "_k": (m, fam, -int(d), var, list(EXIT_LABEL).index(ex), TF_ORDER.index(tf))})
    return pd.DataFrame(rows).sort_values("_k", kind="stable").reset_index(drop=True)


def a_vs_b(tr: pd.DataFrame) -> pd.DataFrame:
    rows = []
    ihs = tr[(tr["family"] == "ihs") & ~tr["is_open"].astype(bool)]
    for (m, ex), g in ihs.groupby(["market", "exit"]):
        for tf in TF_ORDER:
            sub = g if tf == "全部" else g[g["timeframe"] == tf]
            a, b = sub[sub["mode"] == "A"], sub[sub["mode"] == "B"]
            wa, wb = int(a["tp_hit"].astype(bool).sum()), int(b["tp_hit"].astype(bool).sum())
            rows.append({
                "市場": MARKET_LABEL[m], "週期": tf, "出場策略": EXIT_LABEL[ex],
                "A 筆數": len(a), "A 勝率%": 100 * wa / len(a) if len(a) else None,
                "A 平均報酬%": 100 * a["net_ret"].mean() if len(a) else None,
                "A 平均R": _r(a).mean() if len(a) else None,
                "B 筆數": len(b), "B 勝率%": 100 * wb / len(b) if len(b) else None,
                "B 平均報酬%": 100 * b["net_ret"].mean() if len(b) else None,
                "B 平均R": _r(b).mean() if len(b) else None,
                "勝率差（B−A，百分點）": (100 * (wb / len(b) - wa / len(a))) if len(a) and len(b) else None,
                "平均報酬差（B−A，百分點）": (100 * (b["net_ret"].mean() - a["net_ret"].mean()))
                if len(a) and len(b) else None,
                "勝率差 p 值（雙比例 z）": _p_two_prop(wa, len(a), wb, len(b)),
                "報酬差 p 值（常態近似）": _p_welch(a["net_ret"], b["net_ret"]),
                "樣本不足": "是" if min(len(a), len(b)) < MIN_SAMPLE else "",
            })
    return pd.DataFrame(rows)


def tri_stop_compare(tr: pd.DataFrame) -> pd.DataFrame:
    rows = []
    t = tr[(tr["family"] == "tri") & ~tr["is_open"].astype(bool)]
    for (m, d, ex), g in t.groupby(["market", "direction", "exit"]):
        for tf in TF_ORDER:
            sub = g if tf == "全部" else g[g["timeframe"] == tf]
            a, b = sub[sub["stop_kind"] == "inval"], sub[sub["stop_kind"] == "atr"]
            wa, wb = int(a["tp_hit"].astype(bool).sum()), int(b["tp_hit"].astype(bool).sum())
            rows.append({
                "市場": MARKET_LABEL[m], "方向": DIR_LABEL[int(d)], "週期": tf, "出場策略": EXIT_LABEL[ex],
                "失效價 筆數": len(a), "失效價 勝率%": 100 * wa / len(a) if len(a) else None,
                "失效價 平均報酬%": 100 * a["net_ret"].mean() if len(a) else None,
                "失效價 平均R": _r(a).mean() if len(a) else None,
                "失效價 平均停損距離%": 100 * a["risk_pct"].mean() if len(a) else None,
                "ATR 筆數": len(b), "ATR 勝率%": 100 * wb / len(b) if len(b) else None,
                "ATR 平均報酬%": 100 * b["net_ret"].mean() if len(b) else None,
                "ATR 平均R": _r(b).mean() if len(b) else None,
                "ATR 平均停損距離%": 100 * b["risk_pct"].mean() if len(b) else None,
                "平均報酬差（ATR−失效價，百分點）": (100 * (b["net_ret"].mean() - a["net_ret"].mean()))
                if len(a) and len(b) else None,
                "報酬差 p 值（常態近似）": _p_welch(a["net_ret"], b["net_ret"]),
                "樣本不足": "是" if min(len(a), len(b)) < MIN_SAMPLE else "",
            })
    return pd.DataFrame(rows)


def tri_types(tr: pd.DataFrame) -> pd.DataFrame:
    rows = []
    t = tr[(tr["family"] == "tri") & ~tr["is_open"].astype(bool)]
    label = {"ascending": "上升", "descending": "下降", "symmetrical": "對稱"}
    for (m, d, var, ex, ty), g in t.groupby(["market", "direction", "stop_kind", "exit", "sig_triangle_type"]):
        rows.append({"市場": MARKET_LABEL[m], "方向": DIR_LABEL[int(d)], "停損": STOP_LABEL[var],
                     "出場策略": EXIT_LABEL[ex], "三角類型": label.get(ty, ty), **stats(g)})
    return pd.DataFrame(rows)


def yearly(tr: pd.DataFrame) -> pd.DataFrame:
    rows = []
    closed = tr[~tr["is_open"].astype(bool)]
    for (m, fam, d, var, ex, y), g in closed.groupby(["market", "family", "direction", "variant", "exit", "year"]):
        rows.append({"市場": MARKET_LABEL[m], "型態": FAMILY_LABEL[fam], "方向": DIR_LABEL[int(d)],
                     "進場／停損": MODE_LABEL.get(var) if fam == "ihs" else STOP_LABEL[var],
                     "出場策略": EXIT_LABEL[ex], "進場年度": int(y), "筆數": len(g),
                     "勝率%": 100 * g["tp_hit"].astype(bool).mean(), "平均報酬%": 100 * g["net_ret"].mean(),
                     "中位數報酬%": 100 * g["net_ret"].median(), "平均R": _r(g).mean()})
    return pd.DataFrame(rows)


def data_ranges(done: list[dict]) -> pd.DataFrame:
    rows = []
    for m in sorted({d["market"] for d in done}):
        for tf in E.TIMEFRAMES:
            b = [d["bars"][tf] for d in done if d["market"] == m and tf in d.get("bars", {})]
            if not b:
                continue
            n = np.array([x["bars"] for x in b])
            rows.append({"市場": MARKET_LABEL[m], "週期": tf, "序列數": len(b), "K 線總根數": int(n.sum()),
                         "每序列根數 中位數": float(np.median(n)), "最少": int(n.min()), "最多": int(n.max()),
                         "最早 K 線": min(x["first"] for x in b)[:10], "最新 K 線": max(x["last"] for x in b)[:10]})
        srcs = pd.Series([d["source"] for d in done if d["market"] == m]).value_counts()
        rows.append({"市場": MARKET_LABEL[m], "週期": "來源", "序列數": None,
                     "K 線總根數": None, "每序列根數 中位數": None, "最少": None, "最多": None,
                     "最早 K 線": "; ".join(f"{k} {v}" for k, v in srcs.items()), "最新 K 線": None})
    return pd.DataFrame(rows)


def skip_stats(sk: pd.DataFrame) -> pd.DataFrame:
    if not len(sk):
        return pd.DataFrame()
    sk = sk.copy()
    sk["variant"] = np.where(sk["family"] == "ihs", sk["mode"].fillna(""), sk["stop_kind"])
    g = sk.groupby(["market", "family", "direction", "variant", "exit", "reason"]).size().reset_index(name="筆數")
    g["市場"] = g["market"].map(MARKET_LABEL)
    g["型態"] = g["family"].map(FAMILY_LABEL)
    g["方向"] = g["direction"].astype(int).map(DIR_LABEL)
    g["進場／停損"] = [MODE_LABEL.get(v) if f == "ihs" else STOP_LABEL.get(v, v) for f, v in zip(g["family"], g["variant"])]
    g["出場策略"] = g["exit"].map(EXIT_LABEL)
    g["原因"] = g["reason"].map({"position_open": "已有持倉（同 symbol＋週期＋策略）",
                                "stop_beyond_entry": "停損價不在進場價下方（做空：上方）",
                                "target_reached_at_entry": "進場價已達停利價", "missing_level": "缺停損／目標價",
                                "no_atr": "ATR 不足 14 根", "stop_not_positive": "停損價 ≤ 0"}).fillna(g["reason"])
    return g[["市場", "型態", "方向", "進場／停損", "出場策略", "原因", "筆數"]]


NOTES = [
    ("回測編號", "D40（主對話維護 charter；本檔不改 charter）"),
    ("偵測規則", "與線上相同：頭肩底 template profile（含 D33 台股 4h 上限 300、D36）；收斂三角 TRIANGLE_CONFIG（D38，台股 4h 上限 120）。"
                 "ihs_scanner 只 import，不改。"),
    ("逐步前進", "每條序列在 t = 起點, 起點+step, …（step＝pivot 視窗：4h/1d 5 根、1w 3 根）用 df[t−999 : t+1]（1w 為 t−499，"
                 "與線上 --limit 1000、1w 500 根相同）跑偵測。型態以 pattern_id 為鍵，第一次以所需狀態出現的 t 進場。"),
    ("線上彙整", "頭肩底：find_pattern_candidates(include_stale=True) → 去掉 stale → D36 去重 → 分數門檻 → 1d 再對同時間（最近一根已收盤週 K）"
                "的 1w 結果做 D36 跨週期取代。＝ detect_inverse_head_shoulders(include_stale=False)＋build_payload。"
                "三角：detect_triangles(include_stale=False)。stale 型態線上不顯示，所以不當進場訊號。"),
    ("頭肩底進場", "A 突破確認：第一次以 confirmed 出現；B 右肩確認：第一次以 watchlist 出現。進場價＝close[t]。"
                  "停損＝RS 低點；TP1＝對數斐波那契 1.414（D28）；TP2＝突破價×(1+頭到頸線漲幅)，B 用進場當下輸出的預估值，事後不更新。"),
    ("三角進場", "第一次以 breakout_up 出現做多（breakout_down 做空列附錄）。停損 (a) 偵測器 invalidation_price；"
                "(b) 進場價 ∓ 1.5×ATR(14)（Wilder，進場當根已知值）。TP1＝target_conservative、TP2＝target。"
                "三角形結果使用 D41 修正後的 triangles.py（失效價必在突破價另一側，commit b8870ea），所有三角交易同一版程式。"),
    ("出場", "進場後逐根：low ≤ 停損 → 停損（開盤跳空穿過停損則以開盤價成交）；high ≥ TP → 以 TP 成交；同一根兩者都碰 → 算停損。"
            "持有上限 4h 180／1d 120／1w 52 根，到期以收盤出場。資料結束仍未出場者標 open，不列入統計。"),
    ("半半保本", "一半在較近的目標、一半在較遠的目標出場；第一個目標達成後，剩餘部位停損移到進場價（從下一根起算）；"
                "同一根同時達成兩個目標則兩半都出場。勝率＝第一個目標達成比例。"),
    ("成本", "加密每邊 0.1%；台股買進 0.1425%、賣出 0.1425%＋證交稅 0.3%（來回約 0.585%）。所有報酬皆已扣成本。未含滑價。"),
    ("持倉限制", "同一 symbol＋週期＋策略（進場方式／停損方式／出場策略各自獨立）同時最多一筆：訊號 t ≤ 上一筆的出場根即略過（見「略過統計」）。"
                "同一結構的不同變體（不同 LS／RS 組合的 pattern_id）因此不會重疊計數。"),
    ("指標定義", "勝率＝停利出場比例；R＝扣成本報酬 ÷（進場到停損距離／進場價），R 統計不含停損距離 < 0.5% 的交易（約 0.5%，跳空時 R 會失真到 −1000）；平均報酬％即每筆期望值；"
                "獲利因子＝總獲利 ÷ 總虧損；最大連續虧損＝依進場時間排序後連續報酬 < 0 的最多筆數（跨標的）。"),
    ("進場延遲敏感度", "150 個加密＋150 個台股標的的 1d／1w 以 step=1 重跑（docs/backtest-report.md 第 5 節）：A 與三角勝率高約 1–3 個百分點、筆數多約 5–35%，B 無一致差異。"),
    ("偏差與限制", "倖存者偏差（宇宙＝目前上市／交易中的標的，下市的不在內）；進場延遲最多 step−1 根（線上每 4 小時掃描，"
                  "回測每 step 根才看一次）；未含滑價、流動性、台股漲跌停無法成交；台股價格未還原權息（除息缺口可能觸發停損）；"
                  "加密 OKX 來源本機連不上，原本取自 OKX 的幣改由 Gate 永續提供或缺漏；多組比較未做多重比較校正，p 值僅供參考；"
                  "樣本 < 30 的組別標示「樣本不足」。"),
]


def _write_sheet(ws, df: pd.DataFrame) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    cols = [c for c in df.columns if not str(c).startswith("_")]
    ws.append(cols)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in df[cols].itertuples(index=False):
        vals = []
        for v in row:
            if v is None or (isinstance(v, float) and not math.isfinite(v)):
                vals.append(None)
            elif isinstance(v, (np.floating, float)):
                vals.append(round(float(v), 4))
            elif isinstance(v, np.integer):
                vals.append(int(v))
            else:
                vals.append(v)
        ws.append(vals)
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(9, min(28, len(str(c)) * 1.6))
    ws.freeze_panes = "A2"


def write_excel(path: Path, sheets: dict[str, pd.DataFrame], notes: list[tuple[str, str]]) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    wb = Workbook()
    ws = wb.active
    ws.title = "說明"
    ws.append(["項目", "內容"])
    ws["A1"].font = ws["B1"].font = Font(bold=True)
    for k, v in notes:
        ws.append([k, v])
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 120
    for row in ws.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    for name, df in sheets.items():
        _write_sheet(wb.create_sheet(name), df)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _md(df: pd.DataFrame, cols: list[str]) -> str:
    def f(v):
        if v is None or (isinstance(v, float) and not math.isfinite(v)):
            return ""
        if isinstance(v, (float, np.floating)):
            return f"{v:.2f}"
        return str(v)
    lines = ["| " + " | ".join(cols) + " |", "|" + "---|" * len(cols)]
    for _, r in df[cols].iterrows():
        lines.append("| " + " | ".join(f(r[c]) for c in cols) + " |")
    return "\n".join(lines)


def main(tag: str = "main", steps_note: str = "") -> int:
    t0 = time.time()
    tr, sk, done = load(tag)
    if not len(tr):
        print("no trades")
        return 1
    base = D.RESULTS_DIR / tag
    tr.to_csv(base / "trades_all.csv", index=False, encoding="utf-8")
    summ = summarize(tr)
    summ.drop(columns=["_k"]).to_csv(base / "summary.csv", index=False, encoding="utf-8")
    ab, ts_, tt, yr, rng, sks = a_vs_b(tr), tri_stop_compare(tr), tri_types(tr), yearly(tr), data_ranges(done), skip_stats(sk)
    notes = list(NOTES)
    if steps_note:
        notes.append(("本次參數", steps_note))
    errors = [d for d in done if d.get("error")]
    notes.append(("執行", f"標的 {len(done)} 個（錯誤 {len(errors)}）；交易 {len(tr)} 筆（含 open {int(tr['is_open'].astype(bool).sum())} 筆不列入統計）；"
                         f"產生 {time.strftime('%Y-%m-%d %H:%M')}"))
    out = OUTPUT_DIR / f"回測結果_{time.strftime('%Y%m%d')}.xlsx"
    write_excel(out, {"彙總": summ, "頭肩底A_vs_B": ab, "三角停損比較": ts_, "三角類型": tt, "年度分布": yr,
                      "資料範圍": rng, "略過統計": sks}, notes)
    main_cols = ["市場", "型態", "方向", "進場／停損", "週期", "出場策略", "筆數", "勝率%（停利出場比例）",
                 "平均報酬%（每筆期望值）", "中位數報酬%", "平均R（期望值R）", "平均持有天數", "最大連續虧損（筆）", "樣本不足（<30）"]
    with open(base / "summary.md", "w", encoding="utf-8") as fh:
        fh.write("# summary\n\n" + _md(summ, main_cols) + "\n\n# A vs B\n\n"
                 + _md(ab, list(ab.columns)) + "\n\n# triangle stops\n\n" + _md(ts_, list(ts_.columns))
                 + "\n\n# triangle types\n\n" + _md(tt, ["市場", "方向", "停損", "出場策略", "三角類型", "筆數",
                                                       "勝率%（停利出場比例）", "平均報酬%（每筆期望值）", "平均R（期望值R）"])
                 + "\n\n# ranges\n\n" + _md(rng, list(rng.columns)) + "\n\n# skips\n\n"
                 + (_md(sks, list(sks.columns)) if len(sks) else "") + "\n")
    print(f"wrote {out} and {base / 'summary.md'} in {time.time() - t0:.1f}s; trades={len(tr)}")
    return 0
