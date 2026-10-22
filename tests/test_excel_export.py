"""D39: Taiwan 右肩已確認 Excel (ihs_scanner/excel_export.py) and its hook in build_site_tw. No network."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import build_site_tw
from ihs_scanner import excel_export as xe


def tw_row(code, status="watchlist", score=70.0, tf="1d", **extra):
    r = {"symbol": code, "name": f"名{code}", "exchange": "twse" if code < "5000" else "tpex", "market": "stock",
         "timeframe": tf, "status": status, "pattern_score": score, "shape_score": 12.345,
         "current_close": 101.256, "current_neckline_price": 110.0, "distance_to_neckline": 0.0795,
         "neckline_side": "below", "neckline_change_pct": -0.0123, "tp1": 130.123, "tp2": 140.0,
         "invalidation_price": 80.0, "drop_left_pct": 0.2, "drop_right_pct": 0.19, "head_rise_pct": 0.37,
         "neckline_touches": 1, "LS_time": "2026-05-04T01:00:00Z", "H_time": "2026-07-01T01:00:00Z",
         "N2_time": "2026-08-03T01:00:00Z", "RS_time": "2026-09-01T05:00:00Z",
         "RS_confirm_time": "2026-09-08T05:00:00Z", "tv_symbol": f"TWSE:{code}",
         "pattern_id": f"twse|stock|{code}|{tf}|2026-05-04T01:00:00Z|2026-07-01T01:00:00Z|2026-09-01T05:00:00Z"}
    r.update(extra)
    return r


def payload():
    return {"market": "tw", "generated_at": "2026-10-22T06:25:00Z",
            "results": [tw_row("2330", score=65), tw_row("6488", score=80, tf="4h"), tw_row("1101", "forming"),
                        tw_row("2454", "confirmed"), tw_row("3008", score=80, tf="1w")]}


def test_workbook_reads_back(tmp_path):
    out = tmp_path / "rs_confirmed.xlsx"
    assert xe.write_rs_confirmed(payload(), out) == 3
    wb = load_workbook(out)
    assert wb.sheetnames == ["右肩已確認", "說明"]
    ws = wb["右肩已確認"]
    assert [c.value for c in ws[1]] == xe.HEADERS
    assert ws.max_row == 1 + 3  # = number of watchlist rows
    assert [ws.cell(row=i, column=1).value for i in (2, 3, 4)] == ["3008", "6488", "2330"]  # score, then 1w/1d/4h
    for row in ws.iter_rows():
        for cell in row:
            assert not (isinstance(cell.value, str) and cell.value.startswith("=")), cell.coordinate
            assert cell.data_type != "f"
    # values stored as displayed, times in Taipei
    r = {h: ws.cell(row=4, column=i + 1).value for i, h in enumerate(xe.HEADERS)}
    assert r["目前價"] == 101.26 and r["距頸線"] == -0.0795 and r["TP1"] == 130.12
    assert r["右肩 RS"] == datetime(2026, 9, 1) and r["右肩確認"] == datetime(2026, 9, 8, 13, 0)
    r4h = {h: ws.cell(row=3, column=i + 1).value for i, h in enumerate(xe.HEADERS)}
    assert r4h["右肩 RS"] == datetime(2026, 9, 1, 13, 0)
    assert "2026-10-22 14:25" in wb["說明"]["A2"].value and "共 3 筆" in wb["說明"]["A2"].value


def test_hyperlinks(tmp_path):
    out = tmp_path / "x.xlsx"
    xe.write_rs_confirmed(payload(), out)
    ws = load_workbook(out)["右肩已確認"]
    tv = ws.cell(row=4, column=xe.HEADERS.index("TradingView") + 1)
    site = ws.cell(row=4, column=xe.HEADERS.index("網站") + 1)
    assert tv.value == "開啟圖表"
    assert tv.hyperlink.target == "https://www.tradingview.com/chart/?symbol=TWSE%3A2330&interval=1D"
    assert site.hyperlink.target.startswith("https://chasen8.github.io/head-shoulders-site/?market=tw#twse%7Cstock%7C2330%7C1d%7C")
    tv4h = ws.cell(row=3, column=xe.HEADERS.index("TradingView") + 1)
    assert tv4h.hyperlink.target.endswith("&interval=240")


def test_empty_list_still_writes(tmp_path):
    out = tmp_path / "e.xlsx"
    assert xe.write_rs_confirmed({"results": [], "generated_at": "2026-10-22T06:25:00Z"}, out) == 0
    assert load_workbook(out)["右肩已確認"].max_row == 1


def test_build_hook_never_raises(tmp_path, capsys):
    assert build_site_tw.write_excel(payload(), tmp_path / "rs_confirmed.xlsx") is True
    assert (tmp_path / "rs_confirmed.xlsx").exists()
    bad = {"results": [{"status": "watchlist"}]}  # missing fields -> logged, not raised
    assert build_site_tw.write_excel(bad, tmp_path / "bad.xlsx") is False
    assert "not written" in capsys.readouterr().err
    assert build_site_tw.EXCEL_NAME == "rs_confirmed.xlsx"


def test_site_links_excel_for_tw_only():
    root = Path(__file__).resolve().parent.parent / "site"
    html = (root / "index.html").read_text(encoding="utf-8")
    js = (root / "js" / "app.js").read_text(encoding="utf-8")
    assert 'href="data/tw/rs_confirmed.xlsx"' in html and 'id="dl-xlsx-row" hidden' in html
    assert 'state.market === "tw" && !isTri()' in js
